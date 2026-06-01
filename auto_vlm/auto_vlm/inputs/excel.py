"""Excel Method 1 input adapter."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from auto_vlm.inputs.base import AdapterResult
from auto_vlm.models.cases import EvaluationCase, SamplingRequest
from auto_vlm.utils.errors import ToolError


REQUIRED_COLUMNS = {"case_id"}
KNOWN_COLUMNS = {
    "case_id",
    "video_path",
    "qv_video_path",
    "raw_video_path",
    "sampling_frame",
    "sample_count",
    "json_dir",
    "lidar_overlay_path",
    "lidar_json_dir",
    "lidar_json_path",
    "project_type",
    "start_frame",
    "end_frame",
    "frame_list",
    "focus_feature",
    "frame_metadata",
    "memo",
}


def load_cases(input_path: str | Path) -> AdapterResult:
    path = Path(input_path)
    errors: list[ToolError] = []
    cases: list[EvaluationCase] = []

    if not path.exists():
        return AdapterResult(
            errors=[
                ToolError(
                    code="excel_missing",
                    problem="Excel input file does not exist",
                    location=str(path),
                    cause="The --input path could not be found.",
                    fix="Check the input file path and rerun the command.",
                    batch_status="failed",
                )
            ]
        )

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows())
    if not rows:
        return AdapterResult(
            errors=[
                ToolError(
                    code="excel_empty",
                    problem="Excel input file has no rows",
                    location=str(path),
                    cause="The workbook does not contain a header row.",
                    fix="Add a header row with case_id, video_path, and either frame_list or sampling_frame.",
                    batch_status="failed",
                )
            ]
        )

    headers = [_normalize_header(cell.value) for cell in rows[0]]
    header_set = {header for header in headers if header}
    missing_columns = sorted(REQUIRED_COLUMNS - header_set)
    has_video_column = "video_path" in header_set or "qv_video_path" in header_set
    if not has_video_column:
        missing_columns.append("video_path or qv_video_path")
    if missing_columns:
        return AdapterResult(
            errors=[
                ToolError(
                    code="excel_missing_required_columns",
                    problem="Excel input is missing required columns",
                    location=f"{path}: row 1",
                    cause=f"Missing columns: {', '.join(missing_columns)}",
                    fix="Add required columns: case_id and video_path or qv_video_path. Each row also needs exactly one of frame_list or sampling_frame.",
                    batch_status="failed",
                )
            ]
        )

    for row_number, row in enumerate(rows[1:], start=2):
        row_data = {
            header: _cell_value(header, cell)
            for header, cell in zip(headers, row)
        }
        if _is_empty_row(row_data):
            continue
        try:
            cases.append(_row_to_case(path, row_number, row_data))
        except ValueError as exc:
            errors.append(
                ToolError(
                    code="excel_row_invalid",
                    problem="Excel row is invalid",
                    location=f"{path}: row {row_number}",
                    cause=str(exc),
                    fix="Correct the row values and rerun the batch.",
                    case_id=str(row_data.get("case_id") or "") or None,
                    case_status="adapter_error",
                    batch_status="continue",
                )
            )

    return AdapterResult(cases=cases, errors=errors)


def _row_to_case(input_path: Path, row_number: int, row_data: dict[str, Any]) -> EvaluationCase:
    case_id = _required_cell(row_data, "case_id", row_number)
    qv_video_value = row_data.get("qv_video_path") or row_data.get("video_path")
    if qv_video_value in (None, ""):
        raise ValueError(f"row {row_number} column video_path or qv_video_path is required")
    frame_list = _parse_frame_list(row_data.get("frame_list"))
    sampling_frame_value = row_data.get("sampling_frame")
    if sampling_frame_value in (None, ""):
        sampling_frame_value = row_data.get("sample_count")
    sampling_frame = _parse_optional_int(sampling_frame_value, "sampling_frame")
    if bool(frame_list) == (sampling_frame is not None):
        raise ValueError(f"row {row_number} must provide exactly one of frame_list or sampling_frame")
    start_frame = _parse_optional_int(row_data.get("start_frame"), "start_frame")
    end_frame = _parse_optional_int(row_data.get("end_frame"), "end_frame")

    sampling_request = SamplingRequest(
        frame_list=tuple(frame_list),
        start_frame=start_frame,
        end_frame=end_frame,
        sampling_frame=sampling_frame,
    )

    external_metadata = {
        key: value
        for key, value in row_data.items()
        if key and key not in KNOWN_COLUMNS and value not in (None, "")
    }

    return EvaluationCase(
        case_id=str(case_id),
        video_path=Path(str(qv_video_value)),
        qv_video_path=Path(str(qv_video_value)),
        raw_video_path=Path(str(row_data["raw_video_path"])) if row_data.get("raw_video_path") else None,
        sampling_request=sampling_request,
        input_source_type="excel",
        input_source_path=input_path,
        json_dir=Path(str(row_data["json_dir"])) if row_data.get("json_dir") else None,
        lidar_overlay_path=Path(str(row_data["lidar_overlay_path"])) if row_data.get("lidar_overlay_path") else None,
        lidar_json_dir=Path(str(row_data["lidar_json_dir"])) if row_data.get("lidar_json_dir") else None,
        lidar_json_path=Path(str(row_data["lidar_json_path"])) if row_data.get("lidar_json_path") else None,
        project_type=str(row_data["project_type"]) if row_data.get("project_type") else None,
        focus_feature=str(row_data["focus_feature"]) if row_data.get("focus_feature") else "ALL",
        frame_metadata=_parse_metadata(row_data.get("frame_metadata")),
        memo=str(row_data["memo"]) if row_data.get("memo") else None,
        external_metadata=external_metadata,
    )


def _normalize_header(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _cell_value(header: str, cell: Any) -> Any:
    value = cell.value
    if header == "frame_list" and isinstance(value, (int, float)) and "," in str(cell.number_format):
        if float(value).is_integer():
            return f"{int(value):,}"
    return value


def _is_empty_row(row_data: dict[str, Any]) -> bool:
    return all(value in (None, "") for value in row_data.values())


def _required_cell(row_data: dict[str, Any], column: str, row_number: int) -> Any:
    value = row_data.get(column)
    if value in (None, ""):
        raise ValueError(f"row {row_number} column {column} is required")
    return value


def _parse_int(value: Any, field_name: str) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field_name} must be an integer") from exc
    if parsed < 0:
        raise ValueError(f"{field_name} cannot be negative")
    return parsed


def _parse_optional_int(value: Any, field_name: str) -> int | None:
    if value in (None, ""):
        return None
    return _parse_int(value, field_name)


def _parse_frame_list(value: Any) -> list[int]:
    if value in (None, ""):
        return []
    frames: list[int] = []
    for raw_part in str(value).split(","):
        part = raw_part.strip()
        if not part:
            continue
        frames.append(_parse_int(part, "frame_list"))
    return frames


def _parse_metadata(value: Any) -> dict[str, str]:
    if value in (None, ""):
        return {}
    metadata: dict[str, str] = {}
    for raw_part in str(value).split(";"):
        part = raw_part.strip()
        if not part:
            continue
        if "=" not in part:
            metadata[part] = ""
            continue
        key, parsed_value = part.split("=", 1)
        metadata[key.strip()] = parsed_value.strip()
    return metadata
