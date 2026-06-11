"""result.xlsx writer."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from auto_vlm.models.evidence import FrameEvidencePackage
from auto_vlm.models.results import FrameTestResult, PackageReviewResult
from auto_vlm.reports.korean_text import (
    feature_inference,
    feature_observed_evidence,
    feature_summary,
)
from auto_vlm.utils.errors import ToolError


RESULT_COLUMNS = [
    "case_id",
    "video_path",
    "raw_video_path",
    "qv_video_path",
    "json_dir",
    "project_type",
    "sampled_frame",
    "timestamp_sec",
    "focus_feature",
    "review_mode",
    "review_quality_status",
    "frame_result",
    "judgment",
    "confidence",
    "suspicious_type",
    "review_priority",
    "summary",
    "triggered_issue_types",
    "evaluated_issue_types",
    "observed_evidence",
    "inference",
    "json_summary",
    "evidence_image",
    "context_image",
    "raw_frame_image",
    "raw_context_image",
    "qv_overlay_frame_image",
    "qv_overlay_context_image",
    "json_snippet",
    "need_human_review",
    "tool_status",
    "tool_error_message",
    "review_source",
    "reviewer",
    "reviewed_at",
    "review_provider",
    "review_model",
    "prompt_version",
    "review_artifact_version",
    "package_id",
    "input_source_type",
    "input_source_path",
    "evidence_integrity",
    "warnings",
]

FEATURE_RESULT_COLUMNS = [
    "package_id",
    "case_id",
    "sampled_frame",
    "feature",
    "result",
    "confidence",
    "triggered_issue_types",
    "evaluated_issue_types",
    "summary",
    "observed_evidence",
    "inference",
    "ics_crop_image",
    "bev_crop_image",
    "raw_frame_image",
    "qv_overlay_frame_image",
    "json_snippet",
    "json_summary",
]


def write_result_xlsx(
    output_path: str | Path,
    packages: list[FrameEvidencePackage],
    results: dict[str, PackageReviewResult] | None = None,
    errors: list[ToolError] | None = None,
    review_quality_status: str = "not_run",
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "results"
    sheet.append(RESULT_COLUMNS)
    feature_sheet = workbook.create_sheet("feature_results")
    feature_sheet.append(FEATURE_RESULT_COLUMNS)

    results = results or {}
    for package in packages:
        result = results.get(package.package_id, PackageReviewResult())
        sheet.append(_package_row(package, result, review_quality_status))
        for feature_result in result.feature_results:
            feature_sheet.append(_feature_result_row(package, feature_result))

    for error in errors or []:
        sheet.append(_error_row(error))

    _format_sheet(sheet)
    _format_sheet(feature_sheet)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def _feature_result_row(package: FrameEvidencePackage, feature_result) -> list[object]:
    return [
        package.package_id,
        package.case_id,
        package.sampled_frame,
        feature_result.feature,
        feature_result.result.value,
        feature_result.confidence.value,
        ", ".join(feature_result.triggered_issue_types),
        ", ".join(feature_result.evaluated_issue_types),
        "" if feature_result.result == FrameTestResult.PASS else feature_summary(package, feature_result),
        feature_observed_evidence(package, feature_result),
        feature_inference(package, feature_result),
        _feature_packet_path(package, feature_result.feature, "ics_crop_image"),
        _feature_packet_path(package, feature_result.feature, "bev_crop_image"),
        str(package.raw_frame_image) if package.raw_frame_image else "",
        str(package.qv_overlay_frame_image or package.center_frame_image),
        str(package.json_snippet) if package.json_snippet else "",
        package.json_summary or "",
    ]


def _package_row(
    package: FrameEvidencePackage,
    result: PackageReviewResult,
    review_quality_status: str,
) -> list[object]:
    metadata = package.source_case_metadata
    return [
        package.case_id,
        metadata.get("video_path", ""),
        metadata.get("raw_video_path", ""),
        metadata.get("qv_video_path", ""),
        metadata.get("json_dir", ""),
        package.project_type or "",
        package.sampled_frame,
        package.timestamp_sec,
        package.focus_feature,
        package.review_mode,
        review_quality_status,
        result.frame_result.value,
        result.judgment.value,
        result.confidence.value,
        result.suspicious_type.value,
        result.review_priority.value,
        _package_summary(package, result),
        _join_issue_types(feature.triggered_issue_types for feature in result.feature_results),
        _join_issue_types(feature.evaluated_issue_types for feature in result.feature_results),
        _package_observed_evidence(package, result),
        _package_inference(package, result),
        package.json_summary or "",
        str(package.center_frame_image),
        str(package.context_image) if package.context_image else "",
        str(package.raw_frame_image) if package.raw_frame_image else "",
        str(package.raw_context_image) if package.raw_context_image else "",
        str(package.qv_overlay_frame_image or package.center_frame_image),
        str(package.qv_overlay_context_image) if package.qv_overlay_context_image else "",
        str(package.json_snippet) if package.json_snippet else "",
        result.need_human_review,
        result.tool_status,
        result.tool_error_message,
        result.provenance.source.value,
        result.provenance.reviewer,
        result.provenance.reviewed_at,
        result.provenance.provider,
        result.provenance.model,
        result.provenance.prompt_version,
        result.provenance.artifact_version,
        package.package_id,
        metadata.get("input_source_type", ""),
        metadata.get("input_source_path", ""),
        str(package.evidence_integrity.as_dict()),
        "",
    ]


def _error_row(error: ToolError) -> list[object]:
    row = [""] * len(RESULT_COLUMNS)
    row[_column_index("case_id")] = error.case_id or ""
    row[_column_index("focus_feature")] = "ALL"
    row[_column_index("review_mode")] = "unknown"
    row[_column_index("frame_result")] = "fail"
    row[_column_index("judgment")] = "tool_error"
    row[_column_index("summary")] = error.problem
    row[_column_index("tool_status")] = error.case_status
    row[_column_index("tool_error_message")] = error.as_message()
    row[_column_index("review_source")] = "unknown"
    return row


def _join_issue_types(groups: object) -> str:
    values: list[str] = []
    seen: set[str] = set()
    for group in groups:
        for item in group:
            if item in seen:
                continue
            seen.add(item)
            values.append(item)
    return ", ".join(values)


def _feature_packet_path(package: FrameEvidencePackage, feature: str, field_name: str) -> str:
    for packet in package.feature_evidence_packets:
        if packet.feature == feature:
            path = getattr(packet, field_name)
            return str(path) if path else ""
    return ""


def _package_summary(package: FrameEvidencePackage, result: PackageReviewResult) -> str:
    if not result.feature_results:
        return result.summary
    return " | ".join(
        f"{feature.feature}: {feature_summary(package, feature)}"
        for feature in result.feature_results
        if feature.result.value != "pass" and feature_summary(package, feature)
    )


def _join_feature_text(package: FrameEvidencePackage, result: PackageReviewResult, formatter) -> str:
    if not result.feature_results:
        return ""
    return " | ".join(
        f"{feature.feature}: {formatter(package, feature)}"
        for feature in result.feature_results
        if formatter(package, feature)
    )


def _package_observed_evidence(package: FrameEvidencePackage, result: PackageReviewResult) -> str:
    return _join_feature_text(package, result, feature_observed_evidence) or result.observed_evidence


def _package_inference(package: FrameEvidencePackage, result: PackageReviewResult) -> str:
    return _join_feature_text(package, result, feature_inference) or result.inference


def _column_index(name: str) -> int:
    return RESULT_COLUMNS.index(name)


def _format_sheet(sheet) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="F2F4F7")
    header_font = Font(bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    long_text_columns = {
        "summary",
        "observed_evidence",
        "inference",
        "json_summary",
        "evaluated_issue_types",
        "triggered_issue_types",
        "evidence_integrity",
        "warnings",
    }
    path_columns = {
        "video_path",
        "raw_video_path",
        "qv_video_path",
        "json_dir",
        "evidence_image",
        "context_image",
        "raw_frame_image",
        "raw_context_image",
        "qv_overlay_frame_image",
        "qv_overlay_context_image",
        "json_snippet",
        "input_source_path",
        "ics_crop_image",
        "bev_crop_image",
    }
    for index, header_cell in enumerate(sheet[1], start=1):
        header = str(header_cell.value or "")
        width = 18
        if header in long_text_columns:
            width = 52
        elif header in path_columns:
            width = 38
        sheet.column_dimensions[header_cell.column_letter].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
