from __future__ import annotations

from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from auto_vlm.evidence.builder import build_frame_evidence_package
from auto_vlm.models.cases import EvaluationCase, SamplingRequest
from auto_vlm.models.evidence import FeatureEvidencePacket, VideoMetadata
from auto_vlm.models.results import (
    Confidence,
    FeatureReviewResult,
    FrameTestResult,
    Judgment,
    PackageReviewResult,
    ReviewPriority,
    SuspiciousType,
)
from auto_vlm.reports.excel_report import write_result_xlsx
from auto_vlm.reports.html_report import write_summary_html
from auto_vlm.utils.errors import ToolError


def _package(tmp_path: Path):
    center = tmp_path / "output/cases/CASE_001/frames/frame_00000100.jpg"
    raw = tmp_path / "output/cases/CASE_001/raw_frames/frame_00000100.jpg"
    qv = tmp_path / "output/cases/CASE_001/qv_frames/frame_00000100.jpg"
    json_snippet = tmp_path / "output/cases/CASE_001/json_snippets/frame_00000100.json"
    for path in (center, raw, qv, json_snippet):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"fake")
    case = EvaluationCase(
        case_id="CASE_001",
        video_path=Path("sample.mp4"),
        sampling_request=SamplingRequest(frame_list=(100,), sampling_frame=None),
        input_source_type="excel",
        input_source_path=Path("cases.xlsx"),
        json_dir=Path("json"),
        project_type="APTIV_FVC",
    )
    metadata = VideoMetadata(
        video_path=Path("sample.mp4"),
        frame_count=300,
        fps=30.0,
        width=1920,
        height=1080,
    )
    return build_frame_evidence_package(
        case=case,
        sampled_frame=100,
        video_metadata=metadata,
        center_frame_image=center,
        output_root=tmp_path / "output",
        qv_overlay_frame_image=qv,
        raw_frame_image=raw,
        json_snippet=json_snippet,
        json_summary="objects=2; lanes=1",
        sampling_mode=case.sampling_request.mode.value,
    )


def test_write_result_xlsx_links_evidence(tmp_path):
    package = _package(tmp_path)
    output_path = write_result_xlsx(tmp_path / "output/result.xlsx", [package])

    workbook = load_workbook(output_path)
    assert workbook.sheetnames == ["results", "feature_results"]
    sheet = workbook["results"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: sheet.cell(row=2, column=index + 1).value for index, header in enumerate(headers)}

    assert row["case_id"] == "CASE_001"
    assert row["frame_result"] == "needs_review"
    assert row["judgment"] == "needs_more_evidence"
    assert row["review_source"] == "unknown"
    assert row["evidence_image"].endswith("frame_00000100.jpg")
    assert row["qv_overlay_frame_image"].endswith("frame_00000100.jpg")
    assert sheet.cell(row=2, column=headers.index("summary") + 1).alignment.wrap_text is True


def test_write_result_xlsx_includes_feature_result_evidence_sheet(tmp_path):
    package = _package(tmp_path)
    result = PackageReviewResult(
        feature_results=(
            FeatureReviewResult(
                feature="OD",
                result=FrameTestResult.FAIL,
                confidence=Confidence.HIGH,
                evaluated_issue_types=("DEF-OD-BBOX-FIT",),
                triggered_issue_types=("DEF-OD-BBOX-FIT",),
                summary="OD bbox geometry is wrong",
                observed_evidence="raw frame bus, QV overlay oversized box, JSON object id 60",
                inference="OD fails this frame",
            ),
        )
    )

    output_path = write_result_xlsx(
        tmp_path / "output/result.xlsx",
        [package],
        results={package.package_id: result},
    )

    sheet = load_workbook(output_path)["feature_results"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: sheet.cell(row=2, column=index + 1).value for index, header in enumerate(headers)}

    assert row["package_id"] == package.package_id
    assert row["feature"] == "OD"
    assert row["result"] == "fail"
    assert row["triggered_issue_types"] == "DEF-OD-BBOX-FIT"
    assert "raw_frame_image=" in row["observed_evidence"]
    assert "qv_overlay_frame_image=" in row["observed_evidence"]
    assert "json_snippet=" in row["observed_evidence"]
    assert "fail로 판정합니다" in row["inference"]


def test_reports_expose_feature_crop_artifacts(tmp_path):
    package = _package(tmp_path)
    feature_ics = tmp_path / "output/cases/CASE_001/feature_evidence/CASE_001__frame_00000100/OD__ics.jpg"
    feature_bev = tmp_path / "output/cases/CASE_001/feature_evidence/CASE_001__frame_00000100/OD__bev.jpg"
    for path in (feature_ics, feature_bev):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"fake")
    package = replace(
        package,
        feature_evidence_packets=(
            FeatureEvidencePacket(
                package_id=package.package_id,
                feature="OD",
                ics_crop_image=feature_ics,
                bev_crop_image=feature_bev,
            ),
        ),
    )
    result = PackageReviewResult(
        feature_results=(
            FeatureReviewResult(
                feature="OD",
                result=FrameTestResult.FAIL,
                confidence=Confidence.HIGH,
                evaluated_issue_types=("DEF-OD-BBOX-DUP",),
                triggered_issue_types=("DEF-OD-BBOX-DUP",),
                summary="전방 차량 bbox 중복 이슈",
                observed_evidence="raw/qv/bev/json",
                inference="OD fail",
            ),
        )
    )

    xlsx_path = write_result_xlsx(
        tmp_path / "output/result.xlsx",
        [package],
        results={package.package_id: result},
    )
    sheet = load_workbook(xlsx_path)["feature_results"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: sheet.cell(row=2, column=index + 1).value for index, header in enumerate(headers)}

    assert row["ics_crop_image"].endswith("OD__ics.jpg")
    assert row["bev_crop_image"].endswith("OD__bev.jpg")

    html_path = write_summary_html(
        tmp_path / "output/summary.html",
        [package],
        results={package.package_id: result},
    )
    html = html_path.read_text(encoding="utf-8")
    assert "[Crop 산출물]" in html
    assert "기능 ICS crop" in html
    assert "기능 BEV crop" in html


def test_write_result_xlsx_includes_tool_error_row(tmp_path):
    error = ToolError(
        code="missing_video",
        problem="video path does not exist",
        location="row 2 video_path",
        cause="file not found",
        fix="check the path",
        case_id="CASE_BAD",
    )
    output_path = write_result_xlsx(tmp_path / "output/result.xlsx", [], errors=[error])

    workbook = load_workbook(output_path)
    sheet = workbook["results"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: sheet.cell(row=2, column=index + 1).value for index, header in enumerate(headers)}
    assert row["case_id"] == "CASE_BAD"
    assert row["frame_result"] == "fail"
    assert row["judgment"] == "tool_error"


def test_write_summary_html_links_evidence_and_errors(tmp_path):
    package = _package(tmp_path)
    result = PackageReviewResult(
        judgment=Judgment.POTENTIAL_ISSUE,
        confidence=Confidence.LOW,
        suspicious_type=SuspiciousType.VALUE_JUMP,
        review_priority=ReviewPriority.HIGH,
        summary="distance value jumped",
        observed_evidence="raw and QV disagree near ego path",
        inference="potential distance_or_position issue",
        feature_results=(
            FeatureReviewResult(
                feature="OD",
                result=FrameTestResult.FAIL,
                confidence=Confidence.LOW,
                evaluated_issue_types=("DEF-OD-BBOX-DUP", "DEF-OD-CLASS"),
                triggered_issue_types=("DEF-OD-BBOX-DUP",),
                summary="OD distance value jumped",
                observed_evidence="OD object position differs between raw/QV/JSON",
                inference="OD feature fails this frame",
            ),
            FeatureReviewResult(
                feature="LD",
                result=FrameTestResult.PASS,
                confidence=Confidence.MEDIUM,
                evaluated_issue_types=("DEF-LD-RBD-FN", "DEF-LD-RBD-FP"),
                triggered_issue_types=(),
                summary="LD lane evidence is acceptable",
                observed_evidence="lane overlay follows visible lane markings",
                inference="LD passes this frame",
            ),
        ),
    )
    error = ToolError(
        code="missing_video",
        problem="video path does not exist",
        location="row 2 video_path",
        cause="file not found",
        fix="check the path",
        case_id="CASE_BAD",
    )

    output_path = write_summary_html(
        tmp_path / "output/summary.html",
        [package],
        results={package.package_id: result},
        errors=[error],
    )
    html = output_path.read_text(encoding="utf-8")

    assert "CASE_001 · Frame 100" in html
    assert "FAIL" in html
    assert "PASS" in html
    assert "OD" in html
    assert "LD" in html
    assert "DEF-OD-BBOX-DUP" in html
    assert "기능 검토 리포트" in html
    assert "판단 근거" in html
    assert "fail로 판정합니다" in html
    assert "pass로 판정합니다" in html
    assert "QV 오버레이" in html
    assert "원본 프레임" in html
    assert "QV 프레임" in html
    assert "JSON" in html
    assert "raw_frames/frame_00000100.jpg" in html
    assert "qv_frames/frame_00000100.jpg" in html
    assert "json_snippets/frame_00000100.json" in html
    assert "CASE_BAD" in html


def test_write_summary_html_links_evidence_when_vlm_result_missing(tmp_path):
    package = _package(tmp_path)

    output_path = write_summary_html(tmp_path / "output/summary.html", [package])
    html = output_path.read_text(encoding="utf-8")

    assert "CASE_001 · Frame 100" in html
    assert "기능별 검토 결과가 로드되지 않았습니다." in html
    assert "원본 프레임" in html
    assert "QV 프레임" in html
    assert "JSON" in html
    assert "json_snippets/frame_00000100.json" in html
