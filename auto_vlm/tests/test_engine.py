from __future__ import annotations

import json
from pathlib import Path

import cv2
import numpy as np
from openpyxl import Workbook, load_workbook

import auto_vlm.pipeline.engine as engine
from auto_vlm.models.vlm import Feature
from auto_vlm.pipeline.engine import run_excel_batch
from auto_vlm.vlm.reference_context import canonical_issue_type_ids, gtless_single_frame_applicability


def _evaluated_issue_types(feature: str) -> list[str]:
    feature_enum = Feature(feature)
    applicability = gtless_single_frame_applicability(feature_enum)
    return [
        issue_id
        for issue_id in canonical_issue_type_ids(feature_enum)
        if applicability[issue_id] != "not_evaluable"
    ]


def _write_tiny_video(path: Path, frame_count: int = 6, size: tuple[int, int] = (32, 24)) -> None:
    fourcc = cv2.VideoWriter_fourcc(*"mp4v")
    writer = cv2.VideoWriter(str(path), fourcc, 3.0, size)
    try:
        for index in range(frame_count):
            frame = np.full((size[1], size[0], 3), index * 30, dtype=np.uint8)
            writer.write(frame)
    finally:
        writer.release()


def _write_cases(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "case_id",
            "video_path",
            "raw_video_path",
            "sampling_frame",
            "json_dir",
            "project_type",
            "focus_feature",
        ]
    )
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _write_review_results(path: Path, packages) -> None:
    rows = []
    for package in packages:
        package_id = package.package_id
        packets = {packet.feature: packet for packet in package.feature_evidence_packets}
        rows.append(
            {
                "package_id": package_id,
                "feature_results": [
                    {
                        "feature": feature,
                        "result": "pass",
                        "confidence": "medium",
                        "evaluated_issue_types": _evaluated_issue_types(feature),
                        "triggered_issue_types": [],
                        "summary": f"{package_id} {feature} 기준에서 명확한 이슈 없음",
                        "observed_evidence": (
                            f"{package_id} frame 0 raw={packets[feature].raw_frame_image}, "
                            f"QV/ICS={packets[feature].ics_crop_image}, "
                            f"BEV/VCS={packets[feature].bev_crop_image}, "
                            f"JSON={packets[feature].json_snippet} 증거가 일관되며 "
                            "JSON summary objects=1 기준으로 명확한 충돌이 없음"
                        ),
                        "inference": f"{package_id} {feature} 평가 결과 pass",
                        "uncertainty": "단일 프레임 기준 평가",
                    }
                    for feature in packets
                ],
            }
        )
    path.write_text(json.dumps({"results": rows}), encoding="utf-8")


def test_run_excel_batch_creates_reports_and_continues_after_bad_video(tmp_path):
    video_path = tmp_path / "sample.mp4"
    raw_video_path = tmp_path / "sample_raw.mp4"
    json_dir = tmp_path / "json"
    json_dir.mkdir()
    _write_tiny_video(video_path, size=(64, 24))
    _write_tiny_video(raw_video_path, size=(32, 24))
    (json_dir / "00000000.json").write_text(
        json.dumps({"frame_id": 0, "objects": [{"class": "vehicle"}]}),
        encoding="utf-8",
    )
    (json_dir / "00000005.json").write_text(
        json.dumps({"frame_id": 5, "objects": [{"class": "vehicle"}]}),
        encoding="utf-8",
    )

    input_path = tmp_path / "cases.xlsx"
    _write_cases(
        input_path,
        [
            ["CASE_001", str(video_path), str(raw_video_path), 2, str(json_dir), "APTIV_FVC", "OD"],
            ["CASE_BAD", str(tmp_path / "missing.mp4"), "", 2, "", "APTIV_FVC", "ALL"],
        ],
    )

    result = run_excel_batch(input_path, tmp_path / "output")

    assert len(result.packages) == 2
    assert len(result.errors) == 1
    assert result.result_xlsx is None
    assert result.summary_html is None
    assert result.manifest_json is not None and result.manifest_json.exists()
    assert result.review_tasks_json is not None and result.review_tasks_json.exists()
    assert result.packages[0].raw_frame_image is not None
    assert result.packages[0].raw_context_image is None
    assert result.packages[0].qv_overlay_frame_image is not None
    assert result.packages[0].qv_overlay_context_image is None
    assert result.packages[0].evidence_integrity.raw_video_available is True

    manifest = json.loads(result.manifest_json.read_text(encoding="utf-8"))
    assert manifest["manifest_version"] == "run_manifest_v1"
    assert manifest["inputs"]["input_path"] == str(input_path)
    assert manifest["outputs"]["reports"]["result_xlsx"] is None
    assert manifest["outputs"]["reports"]["summary_html"] is None
    assert manifest["outputs"]["review_tasks_json"] == str(result.review_tasks_json)
    assert manifest["counts"]["packages"] == 2
    assert manifest["counts"]["review_results"] == 0
    assert manifest["counts"]["errors"] == 1
    assert manifest["stages"][1]["name"] == "evidence"
    assert manifest["stages"][1]["status"] == "completed"
    assert manifest["stages"][2]["name"] == "review_tasks"
    assert manifest["stages"][2]["status"] == "completed"
    assert manifest["stages"][3]["name"] == "feature_vlm_review"
    assert manifest["stages"][3]["status"] == "pending"
    assert manifest["stages"][4]["name"] == "feature_result_validation"
    assert manifest["stages"][4]["status"] == "not_run"
    assert manifest["stages"][5]["name"] == "cross_feature_audit"
    assert manifest["stages"][5]["status"] == "pending"
    assert manifest["stages"][6]["status"] == "pending_review"
    assert manifest["errors"][0]["code"] == "video_unreadable"

    review_tasks = json.loads(result.review_tasks_json.read_text(encoding="utf-8"))
    assert review_tasks["artifact_version"] == "review_tasks_v1"
    assert review_tasks["counts"]["packages"] == 2
    assert review_tasks["counts"]["feature_tasks"] == 2
    assert review_tasks["model_tasks_root"] == "model/tasks"
    assert review_tasks["packages"][0]["feature_tasks"][0]["required_observed_evidence"]
    assert (tmp_path / "output" / "model" / "tasks" / result.packages[0].package_id / "OD.json").exists()


def test_run_excel_batch_creates_final_reports_only_after_review_results(tmp_path):
    video_path = tmp_path / "sample.mp4"
    raw_video_path = tmp_path / "sample_raw.mp4"
    json_dir = tmp_path / "json"
    output_dir = tmp_path / "output"
    json_dir.mkdir()
    _write_tiny_video(video_path, size=(64, 24))
    _write_tiny_video(raw_video_path, size=(32, 24))
    (json_dir / "00000000.json").write_text(
        json.dumps({"frame_id": 0, "objects": [{"class": "vehicle"}]}),
        encoding="utf-8",
    )
    (json_dir / "00000005.json").write_text(
        json.dumps({"frame_id": 5, "objects": [{"class": "vehicle"}]}),
        encoding="utf-8",
    )

    input_path = tmp_path / "cases.xlsx"
    _write_cases(
        input_path,
        [["CASE_001", str(video_path), str(raw_video_path), 10, str(json_dir), "APTIV_FVC", "OD"]],
    )

    evidence_result = run_excel_batch(input_path, output_dir)
    assert len(evidence_result.packages) == 1
    assert evidence_result.result_xlsx is None
    assert evidence_result.summary_html is None

    review_path = output_dir / "llm_review_results.json"
    _write_review_results(review_path, evidence_result.packages)

    reviewed_result = run_excel_batch(
        input_path,
        output_dir,
        review_results_path=review_path,
        reuse_existing_artifacts=True,
    )

    assert len(reviewed_result.review_results) == 1
    assert reviewed_result.result_xlsx is not None and reviewed_result.result_xlsx.exists()
    assert reviewed_result.summary_html is not None and reviewed_result.summary_html.exists()
    workbook = load_workbook(reviewed_result.result_xlsx)
    sheet = workbook["results"]
    assert sheet.max_row == 2
    assert "CASE_001" in reviewed_result.summary_html.read_text(encoding="utf-8")
    manifest = json.loads(reviewed_result.manifest_json.read_text(encoding="utf-8"))
    assert manifest["outputs"]["reports"]["result_xlsx"] == str(reviewed_result.result_xlsx)
    assert manifest["outputs"]["reports"]["summary_html"] == str(reviewed_result.summary_html)
    assert manifest["stages"][2]["status"] == "completed"
    assert manifest["stages"][3]["status"] == "completed"
    assert manifest["stages"][4]["status"] == "passed"
    assert manifest["stages"][5]["status"] == "accepted"
    assert manifest["stages"][6]["status"] == "completed"
    assert manifest["review_quality"]["status"] == "passed"
    assert manifest["cross_feature_audit"]["accepted"] == 1
    assert (output_dir / "model" / "validation" / reviewed_result.packages[0].package_id / "OD.json").exists()
    assert (output_dir / "model" / "audits" / f"{reviewed_result.packages[0].package_id}.json").exists()


def test_run_excel_batch_still_creates_reports_when_review_quality_fails(tmp_path):
    video_path = tmp_path / "sample.mp4"
    raw_video_path = tmp_path / "sample_raw.mp4"
    json_dir = tmp_path / "json"
    output_dir = tmp_path / "output"
    json_dir.mkdir()
    _write_tiny_video(video_path, size=(64, 24))
    _write_tiny_video(raw_video_path, size=(32, 24))
    (json_dir / "00000000.json").write_text(
        json.dumps({"frame_id": 0, "objects": [{"class": "vehicle"}]}),
        encoding="utf-8",
    )
    (json_dir / "00000005.json").write_text(
        json.dumps({"frame_id": 5, "objects": [{"class": "vehicle"}]}),
        encoding="utf-8",
    )

    input_path = tmp_path / "cases.xlsx"
    _write_cases(
        input_path,
        [["CASE_001", str(video_path), str(raw_video_path), 10, str(json_dir), "APTIV_FVC", "OD"]],
    )

    evidence_result = run_excel_batch(input_path, output_dir)
    package_id = evidence_result.packages[0].package_id
    review_path = output_dir / "llm_review_results.json"
    _write_review_results(review_path, evidence_result.packages)
    data = json.loads(review_path.read_text(encoding="utf-8"))
    data["results"][0]["feature_results"][0]["summary"] = "ÀÁ¿Ã½¾ broken review text"
    review_path.write_text(json.dumps(data), encoding="utf-8")
    (output_dir / "result.xlsx").write_text("stale", encoding="utf-8")
    (output_dir / "summary.html").write_text("stale", encoding="utf-8")

    reviewed_result = run_excel_batch(
        input_path,
        output_dir,
        review_results_path=review_path,
        reuse_existing_artifacts=True,
    )

    assert len(reviewed_result.review_results) == 1
    assert reviewed_result.review_quality.status == "failed"
    assert reviewed_result.result_xlsx is not None and reviewed_result.result_xlsx.exists()
    assert reviewed_result.summary_html is not None and reviewed_result.summary_html.exists()
    manifest = json.loads(reviewed_result.manifest_json.read_text(encoding="utf-8"))
    assert manifest["stages"][4]["status"] == "failed"
    assert manifest["stages"][6]["status"] == "completed"
    assert manifest["review_quality"]["errors"] >= 1


def test_run_excel_batch_reuses_existing_frame_artifacts(tmp_path, monkeypatch):
    video_path = tmp_path / "sample.mp4"
    raw_video_path = tmp_path / "sample_raw.mp4"
    json_dir = tmp_path / "json"
    output_dir = tmp_path / "output"
    json_dir.mkdir()
    _write_tiny_video(video_path, size=(64, 24))
    _write_tiny_video(raw_video_path, size=(32, 24))
    (json_dir / "00000000.json").write_text(
        json.dumps({"frame_id": 0, "objects": [{"class": "vehicle"}]}),
        encoding="utf-8",
    )
    (json_dir / "00000005.json").write_text(
        json.dumps({"frame_id": 5, "objects": [{"class": "vehicle"}]}),
        encoding="utf-8",
    )

    input_path = tmp_path / "cases.xlsx"
    _write_cases(
        input_path,
        [["CASE_001", str(video_path), str(raw_video_path), 10, str(json_dir), "APTIV_FVC", "OD"]],
    )

    first_result = run_excel_batch(input_path, output_dir)
    assert len(first_result.packages) == 1
    assert first_result.generated_artifact_packages == 1
    assert first_result.reused_artifact_packages == 0

    def fail_extract(*args, **kwargs):
        raise AssertionError("frame extraction should not run when artifacts are reused")

    monkeypatch.setattr(engine, "extract_center_frame", fail_extract)

    second_result = run_excel_batch(input_path, output_dir, reuse_existing_artifacts=True)

    assert len(second_result.packages) == 1
    assert second_result.reused_artifact_packages == 1
    assert second_result.generated_artifact_packages == 0
    manifest = json.loads(second_result.manifest_json.read_text(encoding="utf-8"))
    assert manifest["config"]["reuse_existing_artifacts"] is True
    assert manifest["stages"][1]["facts"]["reused_artifact_packages"] == 1
    assert manifest["stages"][1]["facts"]["generated_artifact_packages"] == 0
