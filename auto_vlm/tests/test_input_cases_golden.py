from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from auto_vlm.pipeline.engine import run_excel_batch


EXPECTED_FAIL_ROWS = {
    ("CASE_001__frame_00000520", "LD", "DEF-LD-RBD-LOCALIZATION"),
    ("CASE_002__frame_00000153", "OD", "DEF-OD-BBOX-FIT"),
    ("CASE_002__frame_00000153", "RBD", "DEF-LD-RBD-FN"),
    ("CASE_002__frame_00000235", "OD", "DEF-OD-BBOX-DUP"),
    ("CASE_002__frame_00000235", "RBD", "DEF-LD-RBD-FN"),
}


def test_input_cases_xlsx_golden_fail_rows(tmp_path):
    repo_root = Path(__file__).resolve().parents[1]
    input_path = repo_root / "input_cases.xlsx"
    review_results_path = repo_root / "outputs" / "excel_test_run_latest" / "llm_review_results.json"

    assert input_path.exists(), "input_cases.xlsx is required for the golden workbook regression test"
    assert review_results_path.exists(), (
        "Golden review results are required at "
        "outputs/excel_test_run_latest/llm_review_results.json"
    )

    result = run_excel_batch(
        input_path,
        tmp_path / "input_cases_golden",
        review_results_path=review_results_path,
        reuse_existing_artifacts=False,
    )

    assert result.errors == []
    assert len(result.packages) == 3
    assert len(result.review_results) == 3
    assert result.result_xlsx is not None and result.result_xlsx.exists()

    workbook = load_workbook(result.result_xlsx, read_only=True, data_only=True)
    sheet = workbook["feature_results"]
    headers = [cell.value for cell in sheet[1]]
    columns = {header: index for index, header in enumerate(headers)}

    actual_fail_rows = {
        (
            row[columns["package_id"]],
            row[columns["feature"]],
            row[columns["triggered_issue_types"]],
        )
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[columns["result"]] == "fail"
    }

    assert actual_fail_rows == EXPECTED_FAIL_ROWS
